#!/usr/bin/env python3
"""
Converte a planilha de ficha técnica do Maluzices para a base do protótipo.

A planilha de origem tem três camadas: uma lista de insumos, uma aba por
produto com a receita, e uma aba CMV com os preços. As três se ligam por
nome digitado, e os nomes nem sempre batem. Este script resolve isso e
gera `prototipo/js/base.js`.

Critério do que entra, decidido com a Pâmela:
  - insumo entra se tem nome, preço maior que zero e unidade conhecida;
  - ficha entra se tem rendimento e ao menos um ingrediente que resolve;
  - nada é estimado. O que falta vira pendência no relatório, não vira dado.

Uso:
    python3 ferramentas/converter-planilha.py CAMINHO_DA_PLANILHA.xlsx
"""
import json
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
SAIDA_JS = RAIZ / 'prototipo' / 'js' / 'base.js'
SAIDA_RELATORIO = RAIZ / 'docs' / 'conversao-relatorio.json'

# A planilha mistura KG, Kg, kg e um KR que é erro de digitação. L e LT são
# litro, que o cálculo trata como quilo. M é metro, e só aparece em papel
# manteiga e alumínio, que são embalagem e saem da lista de insumos.
UNIDADES = {
    'KG': 'kg', 'KR': 'kg',
    'L': 'L', 'LT': 'L',
    'UN': 'un', 'UNI': 'un',
    'M': 'm',
}
UNIDADES_EXCLUIDAS = {'m'}

# Categorias deduzidas do nome do produto. É inferência nossa, não dela,
# e está registrada como tal no relatório.
REGRAS_DE_CATEGORIA = [
    ('caldos', ['CALDO', 'CREME DE CASTANHA', 'SOPA']),
    ('paes', ['PAO', 'PÃO', 'FOCACCIA']),
    ('salgados-assados', ['QUICHE', 'EMPADA', 'EMPADINHA', 'TORTA', 'ESFIRA', 'CRACKER', 'GRANOLA']),
    ('salgados-fritos', ['COXINHA', 'CROQUETE', 'NUGGETS', 'BURGER', 'PASTEL']),
    ('sanduiches', ['SANDUBA', 'SANDUICHE', 'WRAP']),
    ('bebidas', ['SUCO', 'LEITE DE', 'BATIDA']),
    ('massas-e-bases', ['MASSA', 'MIX ', 'RECHEIO', 'PURE', 'PURÊ', 'CREME PARA', 'MOLHO', 'MAIONESE',
                        'REQUEIJU', 'MUSSARELA', 'MANTEIGA', 'FRANGO DESFIADO', 'MIX DE SEMENTES']),
]
NOME_DA_CATEGORIA = {
    'caldos': 'Caldos', 'paes': 'Pães', 'salgados-assados': 'Salgados assados',
    'salgados-fritos': 'Salgados fritos', 'sanduiches': 'Sanduíches',
    'bebidas': 'Bebidas', 'massas-e-bases': 'Massas e bases', 'outros': 'Outros',
}


def normalizar(texto):
    t = str(texto or '').strip().upper()
    t = ''.join(c for c in unicodedata.normalize('NFD', t) if unicodedata.category(c) != 'Mn')
    return re.sub(r'\s+', ' ', t.replace('-', ' ').replace('/', ' ')).strip()


def apelido(texto, usados):
    base = normalizar(texto).lower()
    base = re.sub(r'[^a-z0-9]+', '-', base).strip('-') or 'item'
    base = base[:40]
    identificador, n = base, 2
    while identificador in usados:
        identificador, n = f'{base}-{n}', n + 1
    usados.add(identificador)
    return identificador


def titulo(texto):
    """A planilha é toda em maiúscula. Na tela isso grita."""
    miudas = {'de', 'da', 'do', 'das', 'dos', 'e', 'com', 'para', 'sem', 'em', 'a', 'o'}
    palavras = str(texto or '').strip().lower().split()
    saida = []
    for i, p in enumerate(palavras):
        saida.append(p if (i and p in miudas) else p[:1].upper() + p[1:])
    return ' '.join(saida)


def categoria_de(nome):
    alvo = normalizar(nome)
    for chave, termos in REGRAS_DE_CATEGORIA:
        if any(t in alvo for t in termos):
            return chave
    return 'outros'


def ler_planilha(caminho):
    wb = openpyxl.load_workbook(caminho, data_only=True)
    relatorio = {'pendencias': [], 'decisoes': [], 'descartados': []}

    # ---------------------------------------------------------- insumos
    ws = wb['INSUMOS BASE']
    brutos = []
    for r in range(3, ws.max_row + 1):
        nome = ws.cell(row=r, column=3).value
        if not nome or not str(nome).strip():
            continue
        brutos.append({
            'linha': r,
            'nome': str(nome).strip(),
            'unidade_original': str(ws.cell(row=r, column=4).value or '').strip(),
            'pesoBruto': ws.cell(row=r, column=5).value,
            'pesoLiquido': ws.cell(row=r, column=6).value,
            'preco': ws.cell(row=r, column=8).value,
            'fornecedor': str(ws.cell(row=r, column=9).value or '').strip(),
        })

    insumos, vistos, usados_id = [], {}, set()
    embalagens = []
    for b in brutos:
        chave = normalizar(b['nome'])
        unidade = UNIDADES.get(b['unidade_original'].upper())

        if unidade in UNIDADES_EXCLUIDAS:
            embalagens.append({'nome': titulo(b['nome']), 'preco': b['preco'],
                               'unidade': b['unidade_original']})
            relatorio['decisoes'].append(
                f"{b['nome']}: comprado por metro, tratado como embalagem e retirado da lista de insumos.")
            continue

        if chave in vistos:
            relatorio['decisoes'].append(
                f"{b['nome']}: aparece duas vezes na planilha "
                f"(R$ {vistos[chave]['preco']} de {vistos[chave]['fornecedor'] or 'sem fornecedor'} e "
                f"R$ {b['preco']} de {b['fornecedor'] or 'sem fornecedor'}). "
                f"Mantida a primeira, que é a que as fórmulas dela já usavam.")
            continue

        if not isinstance(b['preco'], (int, float)) or b['preco'] <= 0:
            relatorio['descartados'].append(f"insumo {b['nome']}: sem preço na planilha.")
            continue

        if unidade is None:
            relatorio['descartados'].append(
                f"insumo {b['nome']}: unidade '{b['unidade_original']}' não reconhecida.")
            continue

        bruto = b['pesoBruto'] if isinstance(b['pesoBruto'], (int, float)) and b['pesoBruto'] else 1
        liquido = b['pesoLiquido'] if isinstance(b['pesoLiquido'], (int, float)) and b['pesoLiquido'] else bruto

        insumo = {
            'id': apelido(b['nome'], usados_id),
            'nome': titulo(b['nome']),
            'unidade': unidade,
            'pesoBruto': round(bruto, 4),
            'pesoLiquido': round(liquido, 4),
            'preco': round(b['preco'], 4),
            'fornecedor': titulo(b['fornecedor']) if b['fornecedor'] else '',
            # a planilha não guarda data de cotação; fica em branco até a primeira compra lançada
            'cotacao': '',
            'ativo': True,
        }
        if unidade == 'un':
            relatorio['pendencias'].append(
                f"{insumo['nome']}: comprado por unidade. Falta saber quanto pesa uma unidade "
                f"para o custo por quilo fechar.")
        insumos.append(insumo)
        vistos[chave] = b

    # ---------------------------------------------------------- fichas
    abas = [n for n in wb.sheetnames if n not in ('CMV', 'INSUMOS BASE', 'FICHA TECNICA')]
    fichas_brutas = []
    for aba in abas:
        s = wb[aba]

        # As 73 abas foram feitas à mão e o layout escorrega de uma para
        # outra: o rodapé do custo não está sempre na mesma linha. Por isso
        # procuramos pelo rótulo e lemos o primeiro número à direita dele,
        # em vez de confiar na posição.
        rotulos = {}
        linha_ingredientes = None
        for linha in s.iter_rows(min_row=1, max_row=min(s.max_row, 60), max_col=12):
            for celula in linha:
                if not isinstance(celula.value, str):
                    continue
                texto = normalizar(celula.value)
                if texto.startswith('INGREDIENTE') and linha_ingredientes is None:
                    linha_ingredientes = celula.row
                    coluna_ingrediente = celula.column
                for alvo in ('NOME DO PRATO', 'RENDIMENTO DA RECEITA', 'PESO TOTAL',
                             'CUSTO DA RECEITA', 'CUSTO POR PORCAO', 'TEMPO DE PREPARO'):
                    if texto.startswith(alvo) and alvo not in rotulos:
                        valor = next((s.cell(row=celula.row, column=c).value
                                      for c in range(celula.column + 1, min(s.max_column, 12) + 1)
                                      if s.cell(row=celula.row, column=c).value not in (None, '')), None)
                        rotulos[alvo] = valor

        itens = []
        if linha_ingredientes:
            for r in range(linha_ingredientes + 1, linha_ingredientes + 18):
                ing = s.cell(row=r, column=coluna_ingrediente).value
                qtd = s.cell(row=r, column=coluna_ingrediente + 1).value
                if ing and str(ing).strip() and isinstance(qtd, (int, float)) and qtd > 0:
                    itens.append({'nome': str(ing).strip(), 'quantidade': round(qtd, 4)})

        # Em algumas abas a célula ao lado de "NOME DO PRATO" guarda a
        # legenda da foto, não o nome. Nesses casos vale o nome da aba.
        interno = str(rotulos.get('NOME DO PRATO') or '').strip()
        if normalizar(interno) in {'FOTO', 'IMAGEM', 'NOME', 'PRATO'} or len(interno) < 4:
            interno = ''

        fichas_brutas.append({
            'aba': aba,
            'interno': interno,
            'rendimento': rotulos.get('RENDIMENTO DA RECEITA'),
            'itens': itens,
            'pesoTotal': rotulos.get('PESO TOTAL'),
            'custoReceita': rotulos.get('CUSTO DA RECEITA'),
            'custoPorcao': rotulos.get('CUSTO POR PORCAO'),
            'tempoPreparo': rotulos.get('TEMPO DE PREPARO'),
        })

    indice_insumo = {normalizar(i['nome']): i for i in insumos}
    indice_ficha = {}
    for f in fichas_brutas:
        for chave in {normalizar(f['aba']), normalizar(f['interno'])} - {''}:
            indice_ficha.setdefault(chave, f)

    fichas, usados_ficha = [], set()
    for f in fichas_brutas:
        if not f['itens']:
            relatorio['descartados'].append(f"ficha {f['aba']}: sem ingrediente preenchido.")
            continue
        if not isinstance(f['rendimento'], (int, float)) or f['rendimento'] <= 0:
            relatorio['descartados'].append(f"ficha {f['aba']}: sem rendimento.")
            continue

        ingredientes, perdidos = [], []
        for it in f['itens']:
            chave = normalizar(it['nome'])
            if chave in indice_insumo:
                ingredientes.append({'insumoId': indice_insumo[chave]['id'], 'quantidade': it['quantidade']})
            elif chave in indice_ficha:
                ingredientes.append({'subReceitaAba': indice_ficha[chave]['aba'], 'quantidade': it['quantidade']})
            else:
                perdidos.append(it['nome'])

        if not ingredientes:
            relatorio['descartados'].append(f"ficha {f['aba']}: nenhum ingrediente encontrado na lista de insumos.")
            continue
        for p in perdidos:
            relatorio['pendencias'].append(
                f"{titulo(f['aba'])}: o ingrediente '{p}' não existe na lista de insumos e ficou de fora do custo.")

        nome = titulo(f['interno'] or f['aba'])
        if f['interno'] and normalizar(f['interno']) != normalizar(f['aba']):
            relatorio['pendencias'].append(
                f"A aba se chama '{f['aba']}' mas por dentro está escrito '{f['interno']}'. "
                f"Usamos o nome de dentro. Confirmar qual é o certo.")
        fichas.append({
            'id': apelido(f['aba'], usados_ficha),
            'aba': f['aba'],
            'nome': nome,
            'categoriaId': categoria_de(f['aba'] + ' ' + (f['interno'] or '')),
            'rendimento': round(f['rendimento'], 4),
            'ingredientes': ingredientes,
            'tempoPreparo': str(f['tempoPreparo']).strip() if f['tempoPreparo'] else '',
            'custoPorcaoPlanilha': round(f['custoPorcao'], 4) if isinstance(f['custoPorcao'], (int, float)) else None,
            'custoReceitaPlanilha': round(f['custoReceita'], 4) if isinstance(f['custoReceita'], (int, float)) else None,
            'vendavel': False,
            'ativo': True,
        })

    por_aba = {f['aba']: f for f in fichas}
    for f in fichas:
        for ing in f['ingredientes']:
            if 'subReceitaAba' in ing:
                alvo = por_aba.get(ing.pop('subReceitaAba'))
                ing['fichaId'] = alvo['id'] if alvo else None
        f['ingredientes'] = [i for i in f['ingredientes'] if i.get('insumoId') or i.get('fichaId')]

    # --------------------------------- sub-receitas que também são insumo
    # O preço por quilo delas ela digita na lista de insumos. Dá para
    # descobrir quanto sai da receita dividindo o custo pelo preço que ela
    # usa, e é esse número que guardamos. Assim o sistema chega no mesmo
    # preço que ela já pratica, em vez de dividir pelo peso do que entrou,
    # que inclui a água que vai pro ralo.
    # O nome da aba vale mais que o nome escrito dentro dela. A aba
    # "porção REQUEIJU" tem "REQUEIJU" como nome do prato, e sem essa
    # preferência o insumo requeijão apontaria para a ficha da porção, que
    # usa requeijão: a ficha viraria ingrediente de si mesma.
    for preferir_aba in (True, False):
        for f in fichas:
            if f.get('insumoGemeoId'):
                continue
            chave = normalizar(f['aba']) if preferir_aba else normalizar(f['nome'])
            gemeo = next((i for i in insumos
                          if normalizar(i['nome']) == chave and not i.get('fichaId')), None)
            if not gemeo:
                continue
            # a ficha não pode se alimentar do próprio insumo
            if any(ing.get('insumoId') == gemeo['id'] for ing in f['ingredientes']):
                relatorio['pendencias'].append(
                    f"{f['nome']}: a receita lista '{gemeo['nome']}' como ingrediente dela mesma. "
                    f"Mantivemos o preço digitado e não ligamos a ficha, senão o cálculo entra em laço.")
                continue
            f['insumoGemeoId'] = gemeo['id']
            gemeo['fichaId'] = f['id']

    for f in fichas:
        gemeo = next((i for i in insumos if i['id'] == f.get('insumoGemeoId')), None)
        if not gemeo:
            continue
        custo = f.get('custoReceitaPlanilha')
        if isinstance(custo, (int, float)) and custo > 0 and gemeo['preco'] > 0:
            f['rendimentoKg'] = round(custo / gemeo['preco'], 4)
            somado = round(sum(i['quantidade'] for i in f['ingredientes']), 4)
            if abs(f['rendimentoKg'] - somado) / max(somado, 0.001) > 0.05:
                relatorio['decisoes'].append(
                    f"{f['nome']}: a receita soma {somado} kg de ingredientes mas rende "
                    f"{f['rendimentoKg']} kg de produto, segundo o preço de R$ {gemeo['preco']} "
                    f"que ela usa. A diferença é o que se perde no preparo.")

    # ---------------------------------------------------------- preços
    ws = wb['CMV']
    precos = []
    for r in range(2, ws.max_row + 1):
        produto = ws.cell(row=r, column=1).value
        if not produto or not str(produto).strip():
            continue
        precos.append({
            'produto': str(produto).strip(),
            'metaCmv': ws.cell(row=r, column=2).value,
            'custoPorcao': ws.cell(row=r, column=3).value,
            'precoPraticado': ws.cell(row=r, column=5).value,
            'tamanho': ws.cell(row=r, column=6).value,
        })

    for p in precos:
        # casa pelo custo por porção, que é número e não texto. Só depois pelo nome.
        alvo = None
        if isinstance(p['custoPorcao'], (int, float)):
            iguais = [f for f in fichas if f['custoPorcaoPlanilha'] is not None
                      and abs(f['custoPorcaoPlanilha'] - p['custoPorcao']) < 0.005]
            if len(iguais) == 1:
                alvo = iguais[0]
        if alvo is None:
            chave = normalizar(p['produto'])
            candidatos = [f for f in fichas if normalizar(f['aba']) == chave or normalizar(f['nome']) == chave]
            if not candidatos:
                candidatos = [f for f in fichas if chave in normalizar(f['aba']) or normalizar(f['aba']) in chave]
            if len(candidatos) == 1:
                alvo = candidatos[0]
                if isinstance(p['custoPorcao'], (int, float)) and alvo['custoPorcaoPlanilha'] is not None \
                        and abs(alvo['custoPorcaoPlanilha'] - p['custoPorcao']) > 0.005:
                    relatorio['pendencias'].append(
                        f"{p['produto']}: na tabela de preço o custo por porção é "
                        f"R$ {p['custoPorcao']:.2f}, mas a ficha calcula R$ {alvo['custoPorcaoPlanilha']:.2f}. "
                        f"A tabela de preço está defasada em relação à receita.")
        if alvo is None:
            relatorio['pendencias'].append(
                f"{p['produto']}: tem preço na tabela mas não achamos a ficha correspondente.")
            continue

        alvo['vendavel'] = True
        alvo['precoPraticado'] = round(p['precoPraticado'], 2) if isinstance(p['precoPraticado'], (int, float)) else 0
        if isinstance(p['metaCmv'], (int, float)):
            alvo['metaCmv'] = round(p['metaCmv'], 4)
        if p['tamanho']:
            alvo['tamanhoPorcao'] = str(p['tamanho']).strip()

    for f in fichas:
        f.setdefault('precoPraticado', 0)
        f.setdefault('tamanhoPorcao', '')
        f.pop('aba', None)

    # ---------------------------------------------------- cadastros de apoio
    fornecedores, usados_f = [], set()
    for nome in sorted({i['fornecedor'] for i in insumos if i['fornecedor']}):
        fornecedores.append({'id': apelido(nome, usados_f), 'nome': nome, 'telefone': '', 'email': '', 'ativa': True})

    categorias = []
    usadas = {f['categoriaId'] for f in fichas}
    for chave, nome in NOME_DA_CATEGORIA.items():
        if chave in usadas:
            metas = [f['metaCmv'] for f in fichas if f['categoriaId'] == chave and f.get('metaCmv')]
            categorias.append({'id': chave, 'nome': nome,
                               'metaCmv': round(sum(metas) / len(metas), 4) if metas else 0.3, 'ativa': True})

    unidades = [
        {'id': 'quilo', 'nome': 'Quilo', 'sigla': 'kg', 'conversao': 'base do cálculo', 'ativa': True},
        {'id': 'litro', 'nome': 'Litro', 'sigla': 'L', 'conversao': '1 L tratado como 1 kg', 'ativa': True},
        {'id': 'unidade', 'nome': 'Unidade', 'sigla': 'un', 'conversao': 'peso por unidade a informar', 'ativa': True},
    ]

    return {
        'insumos': insumos, 'fichas': fichas, 'categorias': categorias,
        'fornecedores': fornecedores, 'unidades': unidades,
        'embalagens': embalagens, 'relatorio': relatorio,
    }


def gerar_js(base):
    def js(valor):
        return json.dumps(valor, ensure_ascii=False, indent=2)

    return f'''/**
 * Base do Maluzices, convertida da planilha de ficha técnica.
 *
 * ESTE ARQUIVO É GERADO. Não edite à mão: rode
 *   python3 ferramentas/converter-planilha.py <planilha>.xlsx
 *
 * Insumos: {len(base['insumos'])} | Fichas: {len(base['fichas'])} | \
Vendáveis: {sum(1 for f in base['fichas'] if f['vendavel'])}
 */

export const categorias = {js(base['categorias'])};

export const fornecedores = {js(base['fornecedores'])};

export const unidades = {js(base['unidades'])};

export const insumos = {js(base['insumos'])};

export const fichas = {js(base['fichas'])};

/* Comprados por metro e usados como embalagem, fora da lista de insumos. */
export const materiaisDeEmbalagem = {js(base['embalagens'])};
'''


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    base = ler_planilha(sys.argv[1])

    SAIDA_JS.write_text(gerar_js(base), encoding='utf-8')
    SAIDA_RELATORIO.parent.mkdir(exist_ok=True)
    SAIDA_RELATORIO.write_text(json.dumps(base['relatorio'], ensure_ascii=False, indent=2), encoding='utf-8')

    r = base['relatorio']
    print(f"insumos:      {len(base['insumos'])}")
    print(f"fichas:       {len(base['fichas'])}  (vendáveis: {sum(1 for f in base['fichas'] if f['vendavel'])})")
    print(f"categorias:   {len(base['categorias'])}")
    print(f"fornecedores: {len(base['fornecedores'])}")
    print(f"embalagens:   {len(base['embalagens'])}")
    print(f"\ndecisões tomadas: {len(r['decisoes'])}")
    print(f"pendências para a dona: {len(r['pendencias'])}")
    print(f"descartados: {len(r['descartados'])}")
    print(f"\ngerado: {SAIDA_JS.relative_to(RAIZ)}")
    print(f"gerado: {SAIDA_RELATORIO.relative_to(RAIZ)}")


if __name__ == '__main__':
    main()
